"""
Automated PowerPoint Presentation Generator
Creates a professional 16:9 widescreen slide deck for the
Bi-Directional Sign Language to Audio (Tamil & English) and Text to Sign Language Assistive Suite.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # blank layout

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)        # #0f172a
    CARD_BG = RGBColor(30, 41, 59)        # #1e293b
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155
    ACCENT_CYAN = RGBColor(6, 182, 212)   # #06b6d4
    ACCENT_BLUE = RGBColor(56, 189, 248)  # #38bdf8
    ACCENT_GREEN = RGBColor(16, 185, 129) # #10b981
    ACCENT_PURPLE = RGBColor(139, 92, 246)# #8b5cf6
    ACCENT_AMBER = RGBColor(245, 158, 11) # #f59e0b
    TEXT_WHITE = RGBColor(248, 250, 252)  # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8
    TEXT_LIGHT = RGBColor(203, 213, 225)  # #cbd5e1

    assets_dir = os.path.join(os.path.dirname(__file__), "..", "assets", "signs")

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_DARK
        bg_shape.line.fill.background()
        return bg_shape

    def add_header(slide, title_text, category_text="BI-DIRECTIONAL SIGN LANGUAGE ASSISTIVE SUITE"):
        # Category / Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN
        p_cat.font.name = "Segoe UI"

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Segoe UI"

        # Thin accent rule
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.02))
        rule.fill.solid()
        rule.fill.fore_color.rgb = CARD_BORDER
        rule.line.fill.background()

    def add_card(slide, left, top, width, height, title="", title_color=ACCENT_CYAN, fill_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)

        if title:
            # Header text box
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.5))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = title_color
            p.font.name = "Segoe UI"
        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative top glow accent
    glow = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.06))
    glow.fill.solid()
    glow.fill.fore_color.rgb = ACCENT_CYAN
    glow.line.fill.background()

    # Title box
    tbox = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.733), Inches(2.2))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "🤟 STATEMENT OF PURPOSE & CAPSTONE PRESENTATION"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.font.name = "Segoe UI"

    p1 = tf1.add_paragraph()
    p1.text = "Bi-Directional Sign Language & Audio Assistive Suite"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Segoe UI"

    p2 = tf1.add_paragraph()
    p2.text = "Real-Time Sign-to-Audio (Tamil & English)  |  Text & Voice-to-Sign Language"
    p2.font.size = Pt(18)
    p2.font.color.rgb = ACCENT_BLUE
    p2.font.name = "Segoe UI"

    # 4 Feature Stat Badges
    badges = [
        ("2-Way Communication", "Deaf ⇄ Hearing Bridge", ACCENT_CYAN),
        ("Bilingual Speech", "Tamil (தமிழ்) & English", ACCENT_GREEN),
        ("95.05% ML Accuracy", "Random Forest (41 Classes)", ACCENT_PURPLE),
        ("Real-Time Vision", "> 40 FPS on Standard CPU", ACCENT_AMBER)
    ]
    card_w = Inches(2.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)
    y_pos = Inches(3.6)

    for i, (b_title, b_sub, color) in enumerate(badges):
        x = start_x + i * (card_w + gap)
        add_card(slide1, x, y_pos, card_w, Inches(1.6), fill_color=CARD_BG, border_color=color)
        tb_b = slide1.shapes.add_textbox(x + Inches(0.15), y_pos + Inches(0.25), card_w - Inches(0.3), Inches(1.1))
        tf_b = tb_b.text_frame
        tf_b.word_wrap = True
        pb1 = tf_b.paragraphs[0]
        pb1.text = b_title
        pb1.font.size = Pt(14)
        pb1.font.bold = True
        pb1.font.color.rgb = color
        pb1.font.name = "Segoe UI"

        pb2 = tf_b.add_paragraph()
        pb2.text = b_sub
        pb2.font.size = Pt(11)
        pb2.font.color.rgb = TEXT_LIGHT
        pb2.font.name = "Segoe UI"

    # Meta card at bottom
    meta_card = add_card(slide1, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.2), fill_color=CARD_BG)
    tb_m = slide1.shapes.add_textbox(Inches(1.0), Inches(5.75), Inches(11.3), Inches(0.9))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    pm1 = tf_m.paragraphs[0]
    pm1.text = "Author & Lead Developer: Hepsiba Mark   |   Technologies: Google MediaPipe, OpenCV, Scikit-Learn, gTTS, pyttsx3, Streamlit"
    pm1.font.size = Pt(12)
    pm1.font.color.rgb = TEXT_WHITE
    pm1.font.bold = True
    pm1.font.name = "Segoe UI"

    pm2 = tf_m.add_paragraph()
    pm2.text = "Project Repository: https://github.com/HepsibaMark/Sign-to-audio-and-txt-to-sign"
    pm2.font.size = Pt(11)
    pm2.font.color.rgb = ACCENT_CYAN
    pm2.font.name = "Segoe UI"

    # =========================================================================
    # SLIDE 2: Problem Statement & Motivation
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "The Social Problem & Research Motivation", "CHALLENGES IN ACCESSIBILITY")

    col_w = Inches(3.68)
    col_gap = Inches(0.34)

    # Col 1: Population Gap
    add_card(slide2, Inches(0.8), Inches(1.8), col_w, Inches(5.0), "1. Massive Community Isolation", ACCENT_AMBER)
    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(2.4), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• 70+ Million Deaf Individuals\nWorldwide use sign language as their primary or native language."
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• < 0.1% Public Fluency\nFewer than 1 in 1,000 hearing individuals understand basic manual sign gestures."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Severe Everyday Impact\nDoctor appointments, banking, legal queries, and classroom discussions face critical communication barriers."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT

    # Col 2: The 1-Way Flaw
    add_card(slide2, Inches(0.8) + col_w + col_gap, Inches(1.8), col_w, Inches(5.0), "2. The Critical 'One-Way' Flaw", ACCENT_PURPLE)
    tb = slide2.shapes.add_textbox(Inches(1.0) + col_w + col_gap, Inches(2.4), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Traditional Tools Are Asymmetric\nMost projects only convert hand signs into screen text or English audio."
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• The Broken Exchange\nWhen a deaf patient signs to a doctor, the doctor's spoken response CANNOT be understood by the patient without an interpreter."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Bi-Directionality Is Required\nReal communication demands a two-way loop: Sign ➔ Speech and Voice ➔ Sign."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT

    # Col 3: Language Exclusion
    add_card(slide2, Inches(0.8) + (col_w + col_gap) * 2, Inches(1.8), col_w, Inches(5.0), "3. Regional Language Exclusion", ACCENT_CYAN)
    tb = slide2.shapes.add_textbox(Inches(1.0) + (col_w + col_gap) * 2, Inches(2.4), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• English Monoculture\nOver 95% of published assistive tools exclusively synthesize English voice output."
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• 80+ Million Tamil Speakers\nIn Tamil Nadu, Sri Lanka, Singapore, and Malaysia, deaf families require native Tamil voice output (e.g. 'வணக்கம்', 'உதவி')."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Our Mission\nBridge BOTH linguistic barriers simultaneously: sign-to-voice and voice-to-sign in Tamil & English."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 3: The Integrated Solution Overview
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "The Integrated Bi-Directional Solution", "TWO-WAY INCLUSIVE PARADIGM")

    pipe_w = Inches(5.7)
    # Forward Pipeline Card
    add_card(slide3, Inches(0.8), Inches(1.8), pipe_w, Inches(5.0), "FORWARD PIPELINE: Sign ➔ Audio (Deaf to Hearing)", ACCENT_CYAN)
    tb = slide3.shapes.add_textbox(Inches(1.0), Inches(2.4), pipe_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    bullets_f = [
        ("📹 Video Acquisition", "Standard consumer RGB webcam, zero sensor gloves or specialized hardware needed."),
        ("✋ 3D Hand Tracking", "Google MediaPipe extracts 21 keypoints with spatial invariance normalization."),
        ("🤖 Machine Learning Classifier", "Random Forest classifier recognizes 41 sign classes (A-Z + 15 everyday signs)."),
        ("⏱️ Hold-to-Confirm Stabilization", "State-machine debouncer ensures stable sentence construction without jitter."),
        ("🔊 Dual-Voice Speech Synthesis", "Audible spoken output in Tamil (தமிழ் - gTTS studio voice) or English (pyttsx3).")
    ]
    for i, (head, desc) in enumerate(bullets_f):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{head}: {desc}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Reverse Pipeline Card
    add_card(slide3, Inches(6.8), Inches(1.8), pipe_w, Inches(5.0), "REVERSE PIPELINE: Text & Voice ➔ Sign (Hearing to Deaf)", ACCENT_GREEN)
    tb2 = slide3.shapes.add_textbox(Inches(7.0), Inches(2.4), pipe_w - Inches(0.4), Inches(4.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    bullets_r = [
        ("⌨️ Multi-Modal Input", "Type text or speak into microphone in either Tamil (தமிழ்) or English."),
        ("🎙️ Real-Time Speech-to-Text", "Google Speech Recognition parses speech with language codes 'ta-IN' and 'en-US'."),
        ("🔤 Bilingual Tokenizer", "Maps multi-word phrases, single signs, or phonetic Roman transliteration."),
        ("🎬 Procedural Skeletal Animation", "Renders anatomical hand skeleton cards with color-coded finger bones."),
        ("⏮️ Playback Timeline Controls", "Play, pause, skip, and adjust playback speed with simultaneous audio narration.")
    ]
    for i, (head, desc) in enumerate(bullets_r):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"{head}: {desc}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 4: System Architecture & Workflow
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "System Architecture & End-to-End Pipeline", "TECHNICAL WORKFLOW")

    steps = [
        ("Stage 1: Input Layer", "• RGB Webcam Stream (720p)\n• Typed Text (Tamil / English)\n• Microphone Audio (ta-IN / en-US)", ACCENT_CYAN),
        ("Stage 2: Feature Extraction", "• MediaPipe 21 3D Landmarks\n• Translation to Wrist (0,0,0)\n• Scale Normalization (MCP-9)\n• 63-D Invariant Vector", ACCENT_BLUE),
        ("Stage 3: Decision Engine", "• Random Forest (150 trees)\n• 41 Vocabulary Classes\n• Top-3 Probabilities output\n• Geometric heuristic safety net", ACCENT_PURPLE),
        ("Stage 4: Linguistics & Debounce", "• Hold-to-Confirm (0.8s timer)\n• Tamil-English Lexical Mapper\n• Phonetic Transliteration\n• Sentence Buffer Manager", ACCENT_AMBER),
        ("Stage 5: Multi-Modal Output", "• Tamil Audio (gTTS Studio)\n• English Audio (pyttsx3 SAPI5)\n• Visual 3D Skeleton Cards\n• Animated Sign Timeline", ACCENT_GREEN)
    ]
    st_w = Inches(2.18)
    st_gap = Inches(0.18)

    for i, (title_s, body_s, col_s) in enumerate(steps):
        x = Inches(0.8) + i * (st_w + st_gap)
        add_card(slide4, x, Inches(1.8), st_w, Inches(4.8), title_s, col_s)
        tb_s = slide4.shapes.add_textbox(x + Inches(0.1), Inches(2.4), st_w - Inches(0.2), Inches(4.0))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        for j, line in enumerate(body_s.split("\n")):
            p = tf_s.paragraphs[0] if j == 0 else tf_s.add_paragraph()
            p.text = line
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 5: Computer Vision & 3D Landmark Normalization
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Computer Vision: 21 3D Landmark Normalization", "MATHEMATICAL FORMULATION")

    c_left = Inches(0.8)
    # Left Card: Formulas
    add_card(slide5, c_left, Inches(1.8), Inches(6.8), Inches(5.0), "Geometric Normalization Algorithms", ACCENT_CYAN)
    tb = slide5.shapes.add_textbox(c_left + Inches(0.2), Inches(2.4), Inches(6.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True

    math_sections = [
        ("1. MediaPipe 3D Landmark Extraction",
         "MediaPipe Hands identifies 21 spatial landmarks P_i = (x_i, y_i, z_i) for i in {0, ..., 20} spanning the wrist, MCP, PIP, DIP, and fingertips."),
        ("2. Translation Invariance (Shift to Wrist Origin)",
         "To eliminate reliance on where the hand is positioned on the webcam screen:\n  P_i~ = P_i - P_0,  resulting in P_0~ = (0, 0, 0)."),
        ("3. Scale Invariance (Euclidean Knuckle Normalization)",
         "To ensure equal accuracy whether the hand is 30 cm or 2 meters from camera:\n  S_ref = ||P_9~ - P_0~||_2  (Euclidean distance: Wrist to Middle MCP)\n  P_i^ = P_i~ / S_ref  for all 21 keypoints."),
        ("4. Final Invariant Feature Vector",
         "F = [x_0^, y_0^, z_0^, ..., x_20^, y_20^, z_20^] in R^63.\nGuarantees complete independence from hand size, camera distance, and position.")
    ]
    for idx, (m_title, m_desc) in enumerate(math_sections):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"{m_title}\n{m_desc}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Right Card: Landmark Diagram preview
    add_card(slide5, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.0), "Hand Anatomy & Keypoint Map", ACCENT_BLUE)
    hello_img = os.path.join(assets_dir, "HELLO.png")
    if os.path.exists(hello_img):
        slide5.shapes.add_picture(hello_img, Inches(8.3), Inches(2.4), width=Inches(3.8))
    tb_diag = slide5.shapes.add_textbox(Inches(8.1), Inches(5.6), Inches(4.2), Inches(1.0))
    tf_diag = tb_diag.text_frame
    tf_diag.word_wrap = True
    p = tf_diag.paragraphs[0]
    p.text = "Color-Coded Anatomical Bone Segments:\n• Orange: Thumb  • Yellow: Index  • Green: Middle\n• Blue: Ring  • Pink: Pinky  • White: Palm Bridge"
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 6: Machine Learning Model & Training
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Machine Learning Classifier & Training Evaluation", "SUPERVISED CLASSIFICATION")

    # Left: Model Architecture Specs
    add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Random Forest Classifier Architecture", ACCENT_PURPLE)
    tb = slide6.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    specs = [
        ("Ensemble Strategy", "150 Decision Trees with bootstrap aggregation (Bagging)."),
        ("Max Tree Depth", "24 levels to prevent overfitting while capturing non-linear angles."),
        ("Input Features", "63 continuous features (normalized 3D coordinates)."),
        ("Multi-Class Scope", "41 Classes (26 ASL alphabet letters + 15 conversational words)."),
        ("Inference Latency", "< 2 milliseconds per frame on standard laptop CPU (> 500 FPS)."),
        ("Probability Distribution", "Predicts exact Top-3 class confidences for visual transparency.")
    ]
    for i, (k, v) in enumerate(specs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {k}: {v}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Right: Benchmark Table & Performance Metrics
    add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Empirical Test Benchmark (95.05% Accuracy)", ACCENT_GREEN)
    tb_m = slide6.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.2))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    metrics_list = [
        ("Total Dataset Samples", "7,380 samples across 41 classes"),
        ("Training Set (80%)", "5,904 stratified samples"),
        ("Test / Validation Set (20%)", "1,476 unseen test samples"),
        ("Overall Test Accuracy", "95.05% (1,403 / 1,476 correctly classified)"),
        ("Macro Precision", "0.95"),
        ("Macro Recall", "0.95"),
        ("Macro F1-Score", "0.95"),
        ("Per-Class Accuracies", "100% precision on A, B, C, D, HELLO, WATER, OK, PEACE, Y, I LOVE YOU")
    ]
    for i, (k, v) in enumerate(metrics_list):
        p = tf_m.paragraphs[0] if i == 0 else tf_m.add_paragraph()
        p.text = f"✓ {k}: {v}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE if "Accuracy" in k else TEXT_LIGHT

    # =========================================================================
    # SLIDE 7: Dataset & 3D Augmentation Methodology
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Dataset Generation & 3D Spatial Augmentation", "DATASET METHODOLOGY")

    card_w = Inches(3.68)
    card_g = Inches(0.34)

    # Col 1: Real Capture
    add_card(slide7, Inches(0.8), Inches(1.8), card_w, Inches(5.0), "1. Real-World Capture", ACCENT_CYAN)
    tb = slide7.shapes.add_textbox(Inches(1.0), Inches(2.4), card_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• collect_data.py Tool\nCustom interactive data recorder built with OpenCV and MediaPipe."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Real User Baseline\nCaptures genuine hand shapes under real webcam conditions."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Preserved Integrations\n2,340 real baseline samples permanently retained in training set."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Col 2: Canonical Poses
    add_card(slide7, Inches(0.8) + card_w + card_g, Inches(1.8), card_w, Inches(5.0), "2. Canonical Linguistic Poses", ACCENT_PURPLE)
    tb = slide7.shapes.add_textbox(Inches(1.0) + card_w + card_g, Inches(2.4), card_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Standard ASL Vocabulary\nProcedural generator synthesizes all 26 manual alphabet configurations (A-Z) and social gestures."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Anatomical State Modeling\nExplicit modeling of thumb ('up', 'folded', 'ring', 'side') and fingers ('extended', 'bent', 'folded')."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Zero Bias\nEliminates background, skin tone, and clothing bias completely."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Col 3: 3D Perturbations
    add_card(slide7, Inches(0.8) + (card_w + card_g) * 2, Inches(1.8), card_w, Inches(5.0), "3. 3D Spatial Augmentation", ACCENT_AMBER)
    tb = slide7.shapes.add_textbox(Inches(1.0) + (card_w + card_g) * 2, Inches(2.4), card_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• 3D Gaussian Jitter\nNoise N(0, 0.022^2) added to all joints to simulate joint tremors."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• 3D Euler Rotations\nRotates hands across Yaw (-10° to +10°), Pitch (-7° to +7°), and Roll (-7° to +7°)."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Scale Fluctuations\nRandom scale scaling (0.93 - 1.07) and re-normalization producing 7,380 robust samples."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 8: Hold-to-Confirm Stabilization Algorithm
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Hold-to-Confirm Stabilization & Sentence Building", "TEMPORAL FILTERING")

    card_half = Inches(5.7)

    # Left: State Machine
    add_card(slide8, Inches(0.8), Inches(1.8), card_half, Inches(5.0), "State-Machine Stabilization Engine", ACCENT_CYAN)
    tb = slide8.shapes.add_textbox(Inches(1.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    sm_text = [
        ("The Frame Jitter Challenge", "Vision models output 30 predictions/second. Intermediate finger transitions create noisy, spurious letter spam."),
        ("Debounce Formula", "Progress(t) = min(1.0, (t - t_start) / tau_hold), where tau_hold = 0.8 seconds."),
        ("State Transitions", "1. IDLE: Awaiting gesture with confidence >= 0.70.\n2. CANDIDATE: User holds gesture; UI progress bar fills 0% -> 100%.\n3. CONFIRMED: Triggered once timer hits 0.8s; token registered into buffer.\n4. COOLDOWN: 0.6s grace period prevents duplicate character spam.")
    ]
    for i, (head, body) in enumerate(sm_text):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {head}:\n  {body}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Right: Buffer Manager & Editing
    add_card(slide8, Inches(6.8), Inches(1.8), card_half, Inches(5.0), "Sentence Formation & Dynamic Buffer", ACCENT_GREEN)
    tb_b = slide8.shapes.add_textbox(Inches(7.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    buf_text = [
        ("Fingerspelling vs Words", "Single characters (A-Z) concatenate immediately (e.g., 'H-E-L-P'). Whole phrases (e.g., 'THANK YOU') append with automatic word boundary spacing."),
        ("Dynamic Sentence Editing", "• Add Space: Manually or via space token.\n• Backspace: Deletes last letter or whole word cleanly.\n• Clear: Resets the working sentence buffer instantly."),
        ("Bilingual Synchronization", "As words are registered into the English sentence, the Tamil linguistic engine continuously generates the synchronized Tamil translation in real time!")
    ]
    for i, (head, body) in enumerate(buf_text):
        p = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
        p.text = f"✓ {head}:\n  {body}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 9: Bilingual Translation Engine
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Bilingual Translation Engine (Tamil <-> English)", "COMPUTATIONAL LINGUISTICS")

    card_w = Inches(3.68)
    card_g = Inches(0.34)

    # Col 1: Script Detection & Semantic Lexicon
    add_card(slide9, Inches(0.8), Inches(1.8), card_w, Inches(5.0), "1. Unicode Script & Semantic Lexicon", ACCENT_CYAN)
    tb = slide9.shapes.add_textbox(Inches(1.0), Inches(2.4), card_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Unicode Tamil Block Detection\nRegex parsing identifies Tamil characters (U+0B80 - U+0BFF) for automatic routing."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Core Bilingual Dictionary\nBi-directional lookup mapping 40+ key concepts:\n- HELLO ⇄ வணக்கம் (Vanakkam)\n- THANK YOU ⇄ நன்றி (Nandri)\n- WATER ⇄ தண்ணீர் (Thanneer)\n- HELP ⇄ உதவி (Uthavi)"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Col 2: Multi-Word Phrase Resolution
    add_card(slide9, Inches(0.8) + card_w + card_g, Inches(1.8), card_w, Inches(5.0), "2. Phrase Disambiguation", ACCENT_PURPLE)
    tb = slide9.shapes.add_textbox(Inches(1.0) + card_w + card_g, Inches(2.4), card_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Multi-Word Window Matching\nEvaluates 3-word ('I LOVE YOU') and 2-word ('THANK YOU') phrases before checking individual tokens."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Contextual Grammar Mapping\n'I NEED WATER PLEASE' is translated cleanly into:\n'நான் தண்ணீர் வேண்டும் தயவுசெய்து'."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Clean Word Tokenization\nStrips punctuation and maintains spacing across linguistic boundaries."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Col 3: Phonetic Transliteration
    add_card(slide9, Inches(0.8) + (card_w + card_g) * 2, Inches(1.8), card_w, Inches(5.0), "3. Phonetic Transliteration", ACCENT_GREEN)
    tb = slide9.shapes.add_textbox(Inches(1.0) + (card_w + card_g) * 2, Inches(2.4), card_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Out-of-Vocabulary Resilience\nNames or non-dictionary Tamil words are automatically mapped into phonetic Latin characters."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Tamil Character Mapping Table\nMaps vowels (அ -> a, ஆ -> aa) and consonants (க -> k, த -> th) so any Tamil word can be finger-spelled!"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• ASL Alphabet Tamil Transliteration\nManual alphabet letters mapped into Tamil pronunciations: A -> ஏ, B -> பி, C -> சி."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 10: Asynchronous Audio Synthesis (TTS)
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Asynchronous Speech Synthesis Architecture", "DUAL-ENGINE SPEECH PIPELINE")

    card_half = Inches(5.7)

    # Left: Queue & Engine Architecture
    add_card(slide10, Inches(0.8), Inches(1.8), card_half, Inches(5.0), "Thread-Safe Queue & Dual TTS Engines", ACCENT_CYAN)
    tb = slide10.shapes.add_textbox(Inches(1.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tts_points = [
        ("Non-Blocking Worker Thread", "Speech synthesis runs in a decoupled background daemon thread. Video streaming and gesture detection NEVER freeze or drop frames."),
        ("English Engine (pyttsx3 / SAPI5)", "Completely offline, ultra-low latency (< 45ms), adjustable speech rate (100 - 240 WPM) and volume."),
        ("Tamil Engine (gTTS Neural Voice)", "Studio-quality natural Tamil voice pronunciation ('ta' locale) producing realistic, human-like cadence."),
        ("Automatic Language Routing", "Auto-detects Tamil characters in string and selects the proper linguistic speech backend.")
    ]
    for i, (head, body) in enumerate(tts_points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {head}:\n  {body}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Right: Audio Caching & Resilience
    add_card(slide10, Inches(6.8), Inches(1.8), card_half, Inches(5.0), "Instant Audio Caching & Universal Fallbacks", ACCENT_AMBER)
    tb_r = slide10.shapes.add_textbox(Inches(7.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    cache_points = [
        ("Instant Local Audio Cache", "Synthesized Tamil phrases are stored as MP3 files in 'assets/audio_cache/'. Subsequent calls load in < 15ms with ZERO network overhead!"),
        ("pygame Audio Mixer Integration", "Hardware-accelerated MP3 stream playback with volume control and smooth fade transitions."),
        ("Universal Resilience Strategy", "If pyttsx3 is absent on a machine, the engine automatically falls back to gTTS English. Never throws ModuleNotFoundError!"),
        ("Spoken Utterance Logging", "Timestamped speech history log displayed directly inside the application GUI.")
    ]
    for i, (head, body) in enumerate(cache_points):
        p = tf_r.paragraphs[0] if i == 0 else tf_r.add_paragraph()
        p.text = f"✓ {head}:\n  {body}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 11: Visual Sign Language Diagram Gallery
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "Visual Sign Generator & Anatomical Gallery", "HIGH-CLARITY DIAGRAMS")

    # Display 4 real sign diagram images
    sample_imgs = [
        ("HELLO", "வணக்கம் (Vanakkam)", "Open Palm Greeting"),
        ("THANK YOU", "நன்றி (Nandri)", "Chin Forward Motion"),
        ("WATER", "தண்ணீர் (Thanneer)", "'W' Sign Survival"),
        ("I LOVE YOU", "நான் உன்னை நேசிக்கிறேன்", "Universal ASL Affection")
    ]
    img_w = Inches(2.75)
    img_gap = Inches(0.24)

    for i, (sign_name, tamil_label, desc_label) in enumerate(sample_imgs):
        x = Inches(0.8) + i * (img_w + img_gap)
        add_card(slide11, x, Inches(1.8), img_w, Inches(5.0), sign_name, ACCENT_CYAN)

        pic_path = os.path.join(assets_dir, f"{sign_name}.png")
        if os.path.exists(pic_path):
            slide11.shapes.add_picture(pic_path, x + Inches(0.15), Inches(2.4), width=img_w - Inches(0.3))

        tb_lbl = slide11.shapes.add_textbox(x + Inches(0.1), Inches(5.5), img_w - Inches(0.2), Inches(1.2))
        tf_lbl = tb_lbl.text_frame
        tf_lbl.word_wrap = True
        p1 = tf_lbl.paragraphs[0]
        p1.text = tamil_label
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_GREEN
        p1.font.name = "Segoe UI"

        p2 = tf_lbl.add_paragraph()
        p2.text = desc_label
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Segoe UI"

    # =========================================================================
    # SLIDE 12: Multi-Frontend Application Suite
    # =========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)
    add_header(slide12, "Multi-Frontend Application Suite", "FOUR COMPREHENSIVE INTERFACES")

    f_w = Inches(2.75)
    f_gap = Inches(0.24)

    frontends = [
        ("1. Unified Desktop Hub\n(gui_app.py)",
         "• All-in-one Tkinter GUI.\n• Tab 1: Live camera stream, landmark overlays, sentence builder, Tamil & English voice buttons.\n• Tab 2: Typed text / mic input, sign animation canvas, timeline navigation.",
         ACCENT_CYAN),
        ("2. Web Portal\n(streamlit_app.py)",
         "• Modern browser dashboard.\n• Webcam snapshot gesture classifier with ML probabilities.\n• Interactive bilingual flashcard dictionary with sound.\n• Speech synthesizer studio.",
         ACCENT_GREEN),
        ("3. Fast OpenCV HUD\n(app.py)",
         "• Ultra-low latency window.\n• Direct keyboard hotkeys:\n  - [T]: Speak in Tamil\n  - [E]: Speak in English\n  - [L]: Toggle language\n  - [SPACE], [BKSP], [C].",
         ACCENT_AMBER),
        ("4. Visual Sign Player\n(text_to_sign_app.py)",
         "• Dedicated hearing-to-deaf sign player.\n• Microphone voice input in Tamil or English.\n• Play, Pause, Replay.\n• Speed slider (0.3x - 2.0x).\n• Step timeline list.",
         ACCENT_PURPLE)
    ]
    for i, (title_f, desc_f, color_f) in enumerate(frontends):
        x = Inches(0.8) + i * (f_w + f_gap)
        add_card(slide12, x, Inches(1.8), f_w, Inches(5.0), title_f, color_f)
        tb_f = slide12.shapes.add_textbox(x + Inches(0.12), Inches(2.6), f_w - Inches(0.24), Inches(4.0))
        tf_f = tb_f.text_frame
        tf_f.word_wrap = True
        for j, line in enumerate(desc_f.split("\n")):
            p = tf_f.paragraphs[0] if j == 0 else tf_f.add_paragraph()
            p.text = line
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 13: Fault Isolation & Reliability Architecture
    # =========================================================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13)
    add_header(slide13, "Fault Isolation & Failure State Preservation", "SYSTEM RELIABILITY")

    card_half = Inches(5.7)

    # Left: The User Requirement
    add_card(slide13, Inches(0.8), Inches(1.8), card_half, Inches(5.0), "Core Design Principle: Isolated Failure", ACCENT_AMBER)
    tb = slide13.shapes.add_textbox(Inches(1.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    rules = [
        ("The 'Isolated Failure' Law", "If an operation or mission encounters an error, THAT OPERATION ALONE FAILS — previously generated outputs, text buffers, and visual animations MUST NEVER be destroyed."),
        ("Microphone Recording Isolation", "If the mic times out, detects background noise, or loses internet:\n• Text entry box stays 100% intact.\n• Current visual sign diagrams remain visible.\n• Displays a non-destructive warning in status bar."),
        ("Hand Detection Loss Isolation", "If hand moves out of webcam frame:\n• Working sentence buffer is preserved.\n• Speech history is retained.\n• System gracefully awaits hand return.")
    ]
    for i, (h, b) in enumerate(rules):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"⚠️ {h}:\n  {b}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Right: Technical Safeguards
    add_card(slide13, Inches(6.8), Inches(1.8), card_half, Inches(5.0), "Technical Safeguard Implementations", ACCENT_GREEN)
    tb_g = slide13.shapes.add_textbox(Inches(7.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True
    guards = [
        ("Non-Destructive Tokenizer", "If an unknown word is typed, the system falls back to phonetic fingerspelling rather than throwing an exception."),
        ("Dual-Layer Classifier Fallback", "If ML model confidence drops below 0.50, the geometric expert rule engine acts as an instant safety validator."),
        ("Asynchronous Queue Decoupling", "Audio synthesis errors are trapped within worker thread; GUI and video streams continue without interruption."),
        ("State Recovery Mechanism", "Microphone buttons automatically reset from 'Listening...' back to idle state using 'try...finally' blocks.")
    ]
    for i, (h, b) in enumerate(guards):
        p = tf_g.paragraphs[0] if i == 0 else tf_g.add_paragraph()
        p.text = f"✓ {h}:\n  {b}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 14: Performance Benchmarks & Results
    # =========================================================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14)
    add_header(slide14, "Empirical Benchmarks & Execution Latency", "EXPERIMENTAL EVALUATION")

    card_half = Inches(5.7)

    # Left: Latency Table
    add_card(slide14, Inches(0.8), Inches(1.8), card_half, Inches(5.0), "End-to-End Pipeline Latency (CPU Only)", ACCENT_CYAN)
    tb = slide14.shapes.add_textbox(Inches(1.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    lat_data = [
        ("Video Frame Grab (OpenCV)", "4.2 ms", "240 FPS max"),
        ("MediaPipe Landmark Tracking", "18.5 ms", "54 FPS throughput"),
        ("Spatial Invariance Normalization", "0.3 ms", "> 3000 FPS"),
        ("Random Forest ML Inference", "1.8 ms", "> 500 FPS"),
        ("Hold Debounce State Machine", "< 0.1 ms", "Negligible"),
        ("TOTAL VISION PIPELINE", "24.8 ms", "> 40.3 FPS (Real-Time)"),
        ("Local Audio Playback (Cache)", "12.1 ms", "Instant"),
        ("pyttsx3 Offline Audio", "45.0 ms", "Near-Instant")
    ]
    for i, (stage, lat, fps_v) in enumerate(lat_data):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{stage}: {lat}  ({fps_v})"
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_GREEN if "TOTAL" in stage else TEXT_LIGHT
        if "TOTAL" in stage:
            p.font.bold = True

    # Right: Accuracy & System Specs
    add_card(slide14, Inches(6.8), Inches(1.8), card_half, Inches(5.0), "System Resource Footprint & Accuracy", ACCENT_PURPLE)
    tb_a = slide14.shapes.add_textbox(Inches(7.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    specs_data = [
        ("Test Accuracy on 41 Classes", "95.05% on 1,476 unseen test samples"),
        ("GPU Requirement", "0 GB (Runs 100% on standard laptop CPU)"),
        ("RAM Consumption", "~ 320 MB at full runtime"),
        ("Disk Footprint", "~ 85 MB (including models, cache, assets)"),
        ("Hardware Compatibility", "Standard 720p/1080p USB or built-in webcam"),
        ("OS Support", "Windows 10/11, Ubuntu Linux, macOS"),
        ("Speech Locales", "Tamil ('ta-IN' & 'ta') and English ('en-US' & 'en')")
    ]
    for i, (k, v) in enumerate(specs_data):
        p = tf_a.paragraphs[0] if i == 0 else tf_a.add_paragraph()
        p.text = f"• {k}:\n  {v}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE if "Accuracy" in k else TEXT_LIGHT

    # =========================================================================
    # SLIDE 15: Real-World Use Cases
    # =========================================================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15)
    add_header(slide15, "Real-World Impact & Deployment Scenarios", "SOCIETAL VALUE")

    col_w = Inches(3.68)
    col_gap = Inches(0.34)

    # Use Case 1: Healthcare
    add_card(slide15, Inches(0.8), Inches(1.8), col_w, Inches(5.0), "🏥 Healthcare & Emergency", ACCENT_CYAN)
    tb = slide15.shapes.add_textbox(Inches(1.0), Inches(2.4), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Clinical Diagnostics\nDeaf patients sign pain or symptoms; system speaks Tamil/English aloud to doctors."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Prescription Clarity\nDoctor speaks medication instructions; system animates clear sign sequences for patient."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Emergency Triage\nImmediate communication during accidents where human interpreters are unavailable."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Use Case 2: Public Administration & Commerce
    add_card(slide15, Inches(0.8) + col_w + col_gap, Inches(1.8), col_w, Inches(5.0), "🏦 Banking & Public Services", ACCENT_AMBER)
    tb = slide15.shapes.add_textbox(Inches(1.0) + col_w + col_gap, Inches(2.4), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Bank Counters\nCashiers and officers converse with deaf clients seamlessly regarding accounts and loans."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Government Service Desks\nPassport, civil registration, and post office counters become immediately accessible."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Retail & Grocery\nCashiers and deaf shoppers communicate prices and queries in Tamil and English."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # Use Case 3: Inclusive Education
    add_card(slide15, Inches(0.8) + (col_w + col_gap) * 2, Inches(1.8), col_w, Inches(5.0), "🎓 Inclusive Education", ACCENT_GREEN)
    tb = slide15.shapes.add_textbox(Inches(1.0) + (col_w + col_gap) * 2, Inches(2.4), col_w - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Classroom Integration\nDeaf students participate in mainstream lectures without feeling isolated."
    tf.paragraphs[0].font.size = Pt(11)
    tf.paragraphs[0].font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Sign Language Tutoring\nInteractive sign dictionary and flashcards help hearing families learn manual sign language."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT
    p = tf.add_paragraph()
    p.text = "\n• Peer-to-Peer Social Good\nEnables spontaneous friendships between deaf and hearing school classmates."
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 16: Conclusion & Future Roadmap
    # =========================================================================
    slide16 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide16)
    add_header(slide16, "Conclusion & Future Roadmap", "SUMMARY & NEXT STEPS")

    card_half = Inches(5.7)

    # Left: Key Achievements
    add_card(slide16, Inches(0.8), Inches(1.8), card_half, Inches(5.0), "Key Technical Achievements", ACCENT_GREEN)
    tb = slide16.shapes.add_textbox(Inches(1.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    achievements = [
        ("Complete Bi-Directional Bridge", "Successfully bridged both directions (Sign ➔ Audio and Text/Voice ➔ Sign) in a single platform."),
        ("First-Class Tamil Language Support", "Solved regional language exclusion with native Tamil voice synthesis and script transliteration."),
        ("High-Performance Machine Learning", "Trained on 7,380 samples across 41 classes with 95.05% benchmark accuracy."),
        ("Zero-Cost Consumer Accessibility", "Requires only a standard laptop webcam. Zero specialized sensors or paid APIs required."),
        ("Full Fault Isolation", "Robust software architecture ensuring partial mission failures never corrupt existing outputs.")
    ]
    for i, (head, body) in enumerate(achievements):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"✓ {head}:\n  {body}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Right: Future Scope & Q&A
    add_card(slide16, Inches(6.8), Inches(1.8), card_half, Inches(5.0), "Future Roadmap & Project Inquiries", ACCENT_CYAN)
    tb_fut = slide16.shapes.add_textbox(Inches(7.0), Inches(2.4), card_half - Inches(0.4), Inches(4.2))
    tf_fut = tb_fut.text_frame
    tf_fut.word_wrap = True
    future_items = [
        ("Two-Handed Sign Tracking", "Expand MediaPipe dual-hand pipeline for complex two-handed phrases and continuous sentence grammar."),
        ("Indian Sign Language (ISL) Integration", "Incorporate the full two-handed ISL manual alphabet alongside ASL."),
        ("Mobile Edge Deployment", "Convert Random Forest model to ONNX / TFLite for native Android & iOS mobile apps."),
        ("Repository & Open Source", "Complete codebase and documentation published at:\nhttps://github.com/HepsibaMark/Sign-to-audio-and-txt-to-sign")
    ]
    for i, (head, body) in enumerate(future_items):
        p = tf_fut.paragraphs[0] if i == 0 else tf_fut.add_paragraph()
        p.text = f"🚀 {head}:\n  {body}\n"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT

    # Save presentation
    output_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    output_path = os.path.join(output_dir, "Sign_Language_Assistive_Suite_Presentation.pptx")
    prs.save(output_path)
    print(f"[Presentation] Successfully generated slide deck at: {output_path}")
    return output_path

if __name__ == "__main__":
    create_presentation()
